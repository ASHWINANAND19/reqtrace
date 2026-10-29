import csv
import io

from docx import Document
from sqlalchemy import select

from app.exporter import _safe_cell
from app.models import AuditLog
from tests.test_projects import create
from tests.test_testcases import approved_project, gen_tests


def docx_text(content: bytes) -> str:
    doc = Document(io.BytesIO(content))
    cells = [cell.text for table in doc.tables for row in table.rows for cell in row.cells]
    return "\n".join([p.text for p in doc.paragraphs] + cells)


def test_srs_docx_matches_screen(client, analyst):  # TC-23
    pid, srs = approved_project(client, analyst)
    response = client.get(f"/api/projects/{pid}/export/srs.docx", headers=analyst)
    assert response.status_code == 200
    disposition = response.headers["content-disposition"]
    assert "attachment" in disposition and ".docx" in disposition
    text = docx_text(response.content)
    assert srs["title"] in text and "1.1 Purpose" in text and "3.2 Functional Requirements" in text
    for r in srs["requirements"]:
        assert r["req_id"] in text and r["description"] in text


def test_specific_draft_version_can_be_exported(client, analyst):
    pid, srs = approved_project(client, analyst)
    response = client.get(f"/api/projects/{pid}/export/srs.docx?version_id={srs['id']}", headers=analyst)
    assert response.status_code == 200 and "SRS-v1" in response.headers["content-disposition"]


def test_test_cases_csv_has_all_rows_and_columns(client, analyst):  # TC-24
    pid, _ = approved_project(client, analyst)
    cases = gen_tests(client, analyst, pid).json()["test_cases"]
    response = client.get(f"/api/projects/{pid}/export/test-cases.csv", headers=analyst)
    rows = list(csv.reader(io.StringIO(response.content.decode("utf-8-sig"))))
    assert rows[0] == ["Test case", "Requirements", "Type", "Title", "Preconditions", "Steps", "Expected result"]
    assert [r[0] for r in rows[1:]] == [tc["tc_id"] for tc in cases]
    assert all(len(r) == 7 for r in rows)


def test_traceability_csv(client, analyst):
    pid, srs = approved_project(client, analyst)
    gen_tests(client, analyst, pid)
    response = client.get(f"/api/projects/{pid}/export/traceability.csv", headers=analyst)
    rows = list(csv.reader(io.StringIO(response.content.decode("utf-8-sig"))))
    assert [r[0] for r in rows[1:]] == [r["req_id"] for r in srs["requirements"]]
    assert all(r[6] == "yes" for r in rows[1:])


def test_test_plan_docx(client, analyst):
    pid, _ = approved_project(client, analyst)
    assert client.get(f"/api/projects/{pid}/export/test-plan.docx", headers=analyst).status_code == 404
    client.post(f"/api/projects/{pid}/test-plan/generate", headers=analyst)
    text = docx_text(client.get(f"/api/projects/{pid}/export/test-plan.docx", headers=analyst).content)
    assert "1. Scope" in text and "7. Exit Criteria" in text


def test_exports_are_audited(client, analyst, db):  # FR-15
    pid, _ = approved_project(client, analyst)
    client.get(f"/api/projects/{pid}/export/srs.docx", headers=analyst)
    assert client.post(f"/api/projects/{pid}/export/diagram-png?kind=class", headers=analyst).status_code == 204
    details = db.scalars(select(AuditLog.detail).where(AuditLog.action == "export")).all()
    assert details == ["DOCX", "PNG"]


def test_no_approved_srs_cannot_export(client, analyst):
    pid = create(client, analyst)["id"]
    assert client.get(f"/api/projects/{pid}/export/srs.docx", headers=analyst).status_code == 404


def test_csv_formula_injection_neutralised():
    assert _safe_cell("=HYPERLINK(\"x\")") == "'=HYPERLINK(\"x\")"
    assert _safe_cell("Normal text") == "Normal text"
