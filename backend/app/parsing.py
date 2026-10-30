"""Validate and extract text from uploaded requirement files (FR-03, SEC-04)."""
import re
from io import BytesIO
from pathlib import Path

from docx import Document
from pypdf import PdfReader

ALLOWED_EXTENSIONS = (".txt", ".docx", ".pdf")
_CONTROL_CHARS = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]")


class UnsupportedFileError(ValueError):
    pass


class FileTooLargeError(ValueError):
    pass


def clean_text(text: str) -> str:
    """Normalise newlines and drop control characters so stored text is plain data."""
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = _CONTROL_CHARS.sub("", text)
    return re.sub(r"\n{3,}", "\n\n", text).strip()


def safe_filename(name: str) -> str:
    base = Path(name.replace("\\", "/")).name
    return re.sub(r"[^\w.\- ]", "_", base)[:200] or "upload"


def _docx_text(data: bytes) -> str:
    doc = Document(BytesIO(data))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def _pdf_text(data: bytes) -> str:
    reader = PdfReader(BytesIO(data))
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def extract_text(filename: str, data: bytes, max_bytes: int) -> str:
    ext = Path(filename).suffix.lower()
    if ext not in ALLOWED_EXTENSIONS:
        raise UnsupportedFileError(f"Unsupported file type '{ext or 'none'}'. Upload a TXT, DOCX or PDF file.")
    if len(data) > max_bytes:
        raise FileTooLargeError(f"The file is larger than {max_bytes // (1024 * 1024)} MB.")
    if not data:
        raise UnsupportedFileError("The file is empty.")

    # Check the content really matches the extension, not just the name
    try:
        if ext == ".txt":
            if b"\x00" in data:
                raise UnsupportedFileError("This does not look like a plain-text file.")
            text = data.decode("utf-8", errors="replace")
        elif ext == ".docx":
            if not data.startswith(b"PK"):
                raise UnsupportedFileError("This does not look like a valid DOCX file.")
            text = _docx_text(data)
        else:
            if not data.startswith(b"%PDF"):
                raise UnsupportedFileError("This does not look like a valid PDF file.")
            text = _pdf_text(data)
    except UnsupportedFileError:
        raise
    except Exception as exc:  # corrupt document
        raise UnsupportedFileError("The file could not be read. It may be damaged.") from exc

    text = clean_text(text)
    if not text:
        raise UnsupportedFileError("No readable text was found in the file.")
    return text
