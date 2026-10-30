"""Build DOCX and CSV exports (FR-08)."""
import csv
import io
from typing import Any

from docx import Document
from docx.shared import Pt

from app.models import SRSVersion, TestCase, TestPlan

CHAPTERS = {"1": "1. Introduction", "2": "2. Overall Description", "3": "3. Specific Requirements"}


def _safe_cell(value: Any) -> str:
    """Neutralise spreadsheet formula injection (SEC-04): cells starting with = + - @ are prefixed."""
    text = str(value)
    return "'" + text if text[:1] in ("=", "+", "-", "@", "\t", "\r") else text


def _csv(header: list[str], rows: list[list[Any]]) -> bytes:
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(header)
    writer.writerows([[_safe_cell(v) for v in row] for row in rows])
    return buffer.getvalue().encode("utf-8-sig")  # BOM so Excel opens UTF-8 correctly


def _docx_bytes(doc: Document) -> bytes:
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def _requirement_table(doc: Document, rows: list) -> None:
    if not rows:
        doc.add_paragraph("None.")
        return
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    for cell, text in zip(table.rows[0].cells, ("ID", "Title", "Requirement", "Priority"), strict=True):
        cell.text = text
        cell.paragraphs[0].runs[0].bold = True
    for r in rows:
        for cell, text in zip(table.add_row().cells, (r.req_id, r.title, r.description, r.priority), strict=True):
            cell.text = text


def srs_docx(version: SRSVersion) -> bytes:
    doc = Document()
    doc.styles["Normal"].font.size = Pt(11)
    doc.add_heading(version.title, level=0)
    status = f"Version {version.version_no} | {version.status.capitalize()}"
    if version.approver and version.approved_at:
        status += f" by {version.approver.name} on {version.approved_at:%d %b %Y %H:%M} UTC"
    doc.add_paragraph(status)
    chapter = ""
    for s in version.sections:
        number = s["number"].split(".")[0]
        if number != chapter:
            chapter = number
            doc.add_heading(CHAPTERS.get(number, number), level=1)
        doc.add_heading(f"{s['number']} {s['title']}", level=2)
        for paragraph in s["content"].split("\n"):
            if paragraph.strip().startswith("- "):
                doc.add_paragraph(paragraph.strip()[2:], style="List Bullet")
            elif paragraph.strip():
                doc.add_paragraph(paragraph.strip())
    doc.add_heading("3.2 Functional Requirements", level=2)
    _requirement_table(doc, [r for r in version.requirements if r.type == "functional"])
    doc.add_heading("3.3 Non-Functional Requirements", level=2)
    _requirement_table(doc, [r for r in version.requirements if r.type != "functional"])
    return _docx_bytes(doc)


def test_plan_docx(plan: TestPlan, srs: SRSVersion) -> bytes:
    doc = Document()
    doc.add_heading(f"Test Plan: {srs.title}", level=0)
    scope = ", ".join(plan.requirement_ids)
    doc.add_paragraph(f"Based on SRS version {srs.version_no}. Requirements in scope: {scope}")
    for i, s in enumerate(plan.sections, 1):
        doc.add_heading(f"{i}. {s['title']}", level=1)
        for paragraph in s["content"].split("\n"):
            if paragraph.strip().startswith("- "):
                doc.add_paragraph(paragraph.strip()[2:], style="List Bullet")
            elif paragraph.strip():
                doc.add_paragraph(paragraph.strip())
    return _docx_bytes(doc)


def test_cases_csv(cases: list[TestCase]) -> bytes:
    rows = [
        [
            tc.tc_id,
            ", ".join(sorted(link.requirement.req_id for link in tc.links)),
            tc.type,
            tc.title,
            tc.preconditions,
            "\n".join(f"{i}. {step}" for i, step in enumerate(tc.steps, 1)),
            tc.expected_result,
        ]
        for tc in cases
    ]
    return _csv(["Test case", "Requirements", "Type", "Title", "Preconditions", "Steps", "Expected result"], rows)


def traceability_csv(matrix_rows: list) -> bytes:
    rows = [
        [r.req_id, r.title, r.type, r.priority, ", ".join(r.test_cases), len(r.test_cases),
         "yes" if r.covered else "NO"]
        for r in matrix_rows
    ]
    return _csv(["Requirement", "Title", "Type", "Priority", "Test cases", "Count", "Covered"], rows)
